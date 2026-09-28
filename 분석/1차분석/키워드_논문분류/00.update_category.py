"""
전체_법률_목록 카테고리 칼럼 초기화 및 재업데이트 스크립트

사용법:
    python update_category.py
"""

import pandas as pd
import openpyxl

LAW_LIST_PATH    = "00.전체_법률_목록(2026.03.16).csv"
MAPPING_TABLE_PATH = "00.법령분야_카테고리_매핑표.xlsx"


def load_mapping(mapping_path: str) -> dict:
    """매핑표에서 법령분야코드 → {카테고리, 분야명_한글} 딕셔너리 생성"""
    wb = openpyxl.load_workbook(mapping_path)
    ws = wb.active

    mapping = {}
    for r in range(2, ws.max_row + 1):
        코드       = ws.cell(r, 1).value
        분야명_한글 = ws.cell(r, 4).value
        카테고리   = ws.cell(r, 6).value
        if 코드 is not None:
            mapping[int(코드)] = {'카테고리': 카테고리, '분야명_한글': 분야명_한글}

    print(f"[매핑표] 총 {len(mapping)}개 항목 로드")
    print(f"  - 카테고리 확정: {sum(1 for v in mapping.values() if v['카테고리'] != '검토필요')}개")
    print(f"  - 검토필요:      {sum(1 for v in mapping.values() if v['카테고리'] == '검토필요')}개")
    return mapping


def update_category(law_list_path: str, mapping: dict) -> pd.DataFrame:
    """카테고리, 분야명_한글 칼럼 초기화 후 매핑표 기준으로 재업데이트"""
    df = pd.read_csv(law_list_path, encoding='utf-8-sig')

    # 1단계: 칼럼 초기화
    for col in ['카테고리', '분야명_한글']:
        if col in df.columns:
            df[col] = None
        
    # 2단계: 매핑표 기준으로 재업데이트
    df['카테고리']   = df['법령분야코드'].map(lambda x: mapping.get(x, {}).get('카테고리'))
    df['분야명_한글'] = df['법령분야코드'].map(lambda x: mapping.get(x, {}).get('분야명_한글'))

    # 추가: 분야명_한글 → 카테고리 순으로 컬럼 정렬
    base_cols = [c for c in df.columns if c not in ['분야명_한글', '카테고리']]
    df = df[base_cols + ['분야명_한글', '카테고리']]

    # 결과 출력
    print(f"\n[결과] 총 {len(df)}개 법률")
    print("\n카테고리별 법률 수:")
    counts = df['카테고리'].value_counts()
    for cat, cnt in counts.items():
        print(f"  {cat:<15} {cnt:>4}개")

    unmapped = df['카테고리'].isna().sum()
    if unmapped > 0:
        print(f"\n  ※ 매핑 안 된 항목: {unmapped}개")

    return df


def main():
    print("=" * 50)
    print("카테고리 칼럼 초기화 및 재업데이트")
    print("=" * 50)

    # 매핑표 로드
    mapping = load_mapping(MAPPING_TABLE_PATH)

    # 카테고리 업데이트
    df = update_category(LAW_LIST_PATH, mapping)

    # 저장
    df.to_csv(LAW_LIST_PATH, index=False, encoding='utf-8-sig')
    print(f"\n[저장] {LAW_LIST_PATH} 저장 완료")


if __name__ == "__main__":
    main()

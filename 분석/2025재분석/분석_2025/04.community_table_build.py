"""
04.community_table_build.py
────────────────────────────────────────────────────────────
키워드 커뮤니티 요약 테이블 생성

입력 (04.keyword_leiden.py 실행 후 생성됨):
  results/04_keyword_network/keyword_communities.csv
  results/04_keyword_network/community_paper_map.csv
  results/04_keyword_network/community_vs_category.csv

출력:
  results/04_keyword_network/community_table.xlsx

커뮤니티 이름·비고는 스크립트 내 COMMUNITY_META 딕셔너리에서 관리.
이름 변경이 필요하면 해당 딕셔너리만 수정하면 됨.
────────────────────────────────────────────────────────────
"""

import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── 경로 ──────────────────────────────────────────────────────
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
NET_DIR  = os.path.join(BASE_DIR, "results", "04_keyword_network")

# ── 커뮤니티 메타 (수동 관리) ─────────────────────────────────
# key: community_id (int)
# value: (커뮤니티명, 비고)
COMMUNITY_META = {
    0:  ("생성형 AI·저작권",       "ChatGPT 등장 이후 급성장"),
    1:  ("빅데이터·개인정보",       "GDPR 시행(2018) 이후 확산"),
    2:  ("알고리즘·머신러닝",       "2016년 최초 등장 커뮤니티"),        # 옛 3
    3:  ("4차 산업혁명·거버넌스",   "가장 이질적 카테고리 혼합"),        # 옛 2
    4:  ("AI 기본법",             "EU AI Act 발의(2021) 이후 급부상"),  # 옛 6
    5:  ("제조물책임·불법행위",     "의료 AI 사고 책임 논의 포함"),      # 옛 4
    6:  ("자율주행·규제",           "드론·도심항공 포함"),              # 옛 5
    7:  ("AI 법인격·형사책임",      "존재론적·철학적 논의 포함"),
    8:  ("딥페이크·표현의 자유",    "허위정보·초상권 교차 영역"),
    9:  ("AI 발명·특허",            "DABUS 판결 관련 논의"),
    10: ("자율운항선박",            "해사법 특화 소규모 클러스터"),
    11: ("리걸테크",                "법률 AI 서비스 규제 논의"),        # 옛 12
    12: ("비교법 (중국)",           "중국 AI 법제 비교 연구"),          # 옛 13
    13: ("부정경쟁방지",            "소규모 위성 클러스터"),            # 옛 11
}

# ── 데이터 로드 ───────────────────────────────────────────────
kw_com    = pd.read_csv(os.path.join(NET_DIR, "keyword_communities.csv"),   encoding="utf-8-sig")
paper_map = pd.read_csv(os.path.join(NET_DIR, "community_paper_map.csv"),   encoding="utf-8-sig")
cvc       = pd.read_csv(os.path.join(NET_DIR, "community_vs_category.csv"), encoding="utf-8-sig")

# ── 커뮤니티별 노드 수 ────────────────────────────────────────
node_count = kw_com.groupby("community_id").size().rename("node_count")

# ── 커뮤니티별 논문 수 ────────────────────────────────────────
paper_count = paper_map.groupby("primary_community").size().rename("paper_count")

# ── 커뮤니티별 대표 키워드 (빈도 상위 8개) ───────────────────
def top_keywords(df, n=8):
    return (
        df.sort_values("freq", ascending=False)
          .head(n)["keyword"]
          .tolist()
    )

rep_kw = (
    kw_com.groupby("community_id")[["freq", "keyword"]]
          .apply(top_keywords)
          .rename("rep_keywords")
)

# ── 커뮤니티별 주요 카테고리 (상위 1~2개) ────────────────────
# cvc: community_label × category, 'All' 컬럼 제외
# community_label → community_id 매핑 먼저
label_to_id = {
    row["community_label"]: cid
    for cid, row in kw_com.drop_duplicates("community_id")
                           .set_index("community_id")
                           .iterrows()
}

def dominant_categories(row, top_n=2):
    """단일 카테고리 컬럼만 필터링해서 상위 n개 반환."""
    cats = row.drop(labels=["community_label", "All"], errors="ignore")
    # 융복합 제외하고 단일만, 숫자형으로 변환
    single = cats[[c for c in cats.index if not c.startswith("융복합")]].astype(float)
    top = single.nlargest(top_n)
    parts = [f"{cat} ({int(cnt)}편)" for cat, cnt in top.items() if cnt > 0]
    return " / ".join(parts) if parts else "—"

cvc["주요 카테고리"] = cvc.apply(dominant_categories, axis=1)
cvc["community_id"] = cvc["community_label"].map(label_to_id)

main_cat = cvc.set_index("community_id")["주요 카테고리"]

# ── 테이블 조립 ───────────────────────────────────────────────
rows = []
for cid in sorted(COMMUNITY_META.keys()):
    name, note = COMMUNITY_META[cid]
    rows.append({
        "#":        cid,
        "커뮤니티명": name,
        "노드\n(키워드 수)": int(node_count.get(cid, 0)),
        "논문 수":   int(paper_count.get(cid, 0)),
        "대표 키워드": ", ".join(rep_kw.get(cid, [])),
        "주요 카테고리": main_cat.get(cid, "—"),
        "비고":      note,
    })

df_table = pd.DataFrame(rows)

# 합계 행
total_row = {
    "#": "합계",
    "커뮤니티명": "",
    "노드\n(키워드 수)": int(node_count.sum()),
    "논문 수": int(paper_count.sum()),
    "대표 키워드": "",
    "주요 카테고리": "",
    "비고": "",
}
df_table = pd.concat([df_table, pd.DataFrame([total_row])], ignore_index=True)

print(df_table.to_string(index=False))

# ── Excel 저장 ────────────────────────────────────────────────
out_path = os.path.join(NET_DIR, "community_table.xlsx")

with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
    # 타이틀 행을 위해 startrow=2
    df_table.to_excel(writer, index=False, sheet_name="커뮤니티", startrow=2)

    ws = writer.sheets["커뮤니티"]

    # ── 타이틀 병합 ──────────────────────────────────────────
    NCOLS = len(df_table.columns)
    title_text = (
        f"키워드 공출현 네트워크 커뮤니티 분석 결과 "
        f"(Leiden, resolution=1.0)  |  "
        f"총 {len(COMMUNITY_META)}개 커뮤니티  |  "
        f"분석 노드 {int(node_count.sum())}개  |  "
        f"Modularity = 0.64  |  "
        f"매핑 논문 {int(paper_count.sum()):,}편"
    )
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=NCOLS)
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font      = Font(bold=True, size=11)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill      = PatternFill("solid", fgColor="1F4E79")
    title_cell.font      = Font(bold=True, size=11, color="FFFFFF")
    ws.row_dimensions[1].height = 22

    # ── 헤더 스타일 ──────────────────────────────────────────
    header_fill = PatternFill("solid", fgColor="2E75B6")
    for col_idx in range(1, NCOLS + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.fill      = header_fill
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[3].height = 30

    # ── 데이터 행 스타일 ─────────────────────────────────────
    alt_fill  = PatternFill("solid", fgColor="EBF3FB")
    last_fill = PatternFill("solid", fgColor="D6E4F0")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    n_data = len(df_table)
    for row_idx in range(4, 4 + n_data):
        is_total = (row_idx == 3 + n_data)
        for col_idx in range(1, NCOLS + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border    = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if is_total:
                cell.fill = last_fill
                cell.font = Font(bold=True, size=10)
            elif (row_idx - 4) % 2 == 0:
                cell.fill = alt_fill
        ws.row_dimensions[row_idx].height = 40 if not is_total else 18

    # ── 열 너비 ──────────────────────────────────────────────
    col_widths = {
        1: 6,    # #
        2: 18,   # 커뮤니티명
        3: 12,   # 노드
        4: 10,   # 논문 수
        5: 52,   # 대표 키워드
        6: 28,   # 주요 카테고리
        7: 30,   # 비고
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 빈 타이틀 행(row=2) 숨기기
    ws.row_dimensions[2].height = 0

print(f"\n✅ 저장 완료 → {out_path}")

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BASE = Path(__file__).parent
OUT = BASE / "4장_분석결과_초안.docx"

TABLE_DIR = BASE / "results" / "paper_tables"
IMG_DIR_DESC = BASE / "results" / "02_descriptive"
IMG_DIR_KEYWORD = BASE / "results" / "05_keyword_trend"
IMG_DIR_COMMUNITY = BASE / "results" / "05_keyword_community_trend"
IMG_DIR_CITATION = BASE / "results" / "03_citation"

DOC_FONT = "Apple SD Gothic Neo"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(8.8)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def style_table(table, header_fill: str = "1D3557") -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    run.font.name = DOC_FONT
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
                    run.font.size = Pt(8.8)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F3F6F8")


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    run = p.add_run(text)
    run.font.name = DOC_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    run.font.size = Pt(9)
    run.bold = True
    run.font.color.rgb = RGBColor(65, 65, 65)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = DOC_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(85, 85, 85)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(0.6)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(7)


def format_cell_value(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def add_csv_table(doc: Document, csv_path: Path, caption: str, columns: list[str] | None = None) -> None:
    df = pd.read_csv(csv_path, encoding="utf-8-sig")
    if columns:
        df = df[columns]
    table = doc.add_table(rows=1, cols=len(df.columns))
    for idx, col in enumerate(df.columns):
        set_cell_text(table.rows[0].cells[idx], col, bold=True, color="FFFFFF")
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for idx, col in enumerate(df.columns):
            set_cell_text(cells[idx], format_cell_value(row[col]))
    style_table(table)
    add_caption(doc, caption)


def add_image(doc: Document, path: Path, caption: str, width: float = 6.6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def set_table_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_community_features_table(doc: Document) -> None:
    df = pd.read_csv(TABLE_DIR / "table_05_community_features.csv", encoding="utf-8-sig")
    df = df.rename(
        columns={
            "커뮤니티명": "연구 주제군",
            "매핑 논문 수": "논문 수",
            "법학적 해석": "주요 쟁점",
        }
    )[["연구 주제군", "논문 수", "대표 키워드", "주요 쟁점"]]

    widths = [1.35, 0.6, 2.4, 2.1]
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    for idx, col in enumerate(df.columns):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, col, bold=True, color="FFFFFF")
        set_table_cell_width(cell, widths[idx])

    for _, row in df.iterrows():
        cells = table.add_row().cells
        for idx, col in enumerate(df.columns):
            set_cell_text(cells[idx], format_cell_value(row[col]))
            set_table_cell_width(cells[idx], widths[idx])

    style_table(table)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 1 or row is table.rows[0] else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(7.2)
    add_caption(doc, "표 4-2. 키워드 공출현 네트워크의 주요 연구 주제군")


def add_two_images(doc: Document, left: Path, right: Path, caption: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, path in zip(table.rows[0].cells, [left, right]):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(3.05))
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                border = OxmlElement(f"w:{border_name}")
                border.set(qn("w:val"), "nil")
                tc_pr.append(border)
    add_caption(doc, caption)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = DOC_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35

    for style_name, size, color in [
        ("Title", 18, "1D3557"),
        ("Heading 1", 15, "1D3557"),
        ("Heading 2", 12.5, "1D3557"),
    ]:
        style = styles[style_name]
        style.font.name = DOC_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
    return doc


def main() -> None:
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("4. 분석 결과")

    add_body(
        doc,
        "이 장은 KCI 등재 법학 분야 AI 관련 논문 2,014편을 대상으로 한 기술통계, 법 분야 분류, 주요 키워드 변화, 키워드 공출현 네트워크 기반 연구 주제군, 주요 주제군의 시계열 변화 및 내부 인용의 축적 양상을 제시한다. 결과 서술은 전체 규모의 변화에서 출발하여 법 분야 분포, 키워드 지형, 연구 주제군, 주요 주제군의 성장 양상, 연구 기반의 형성으로 이어진다.",
    )

    doc.add_heading("4.1 연도별 전체 논문 수", level=2)
    add_body(
        doc,
        "분석 대상 논문은 2016년 47편에서 2025년 457편으로 증가하였다. 특히 2023년 이후 증가세가 뚜렷하며, 2025년 단일 연도 논문 수는 전체 분석 대상의 22.7%를 차지한다.",
    )
    add_csv_table(
        doc,
        TABLE_DIR / "table_01_annual_counts.csv",
        "표 4-1. 연도별 KCI 등재 법학 분야 AI 관련 논문 수",
    )

    doc.add_heading("4.2 법 분야별 분포: 단일/융복합, 주요 카테고리", level=2)
    add_body(
        doc,
        "법 분야 분류 결과 전체 2,014편 중 단일 분야 논문은 1,410편, 융복합 논문은 604편으로 나타났다. 단일 분야에서는 지식재산권법, 인공지능법, 형사법, 데이터법, 공법이 주요 비중을 차지하였다.",
    )
    add_image(
        doc,
        IMG_DIR_DESC / "category_pie_major.png",
        "그림 4-1. 법 분야별 분포: 융복합 및 주요 단일 분야",
        width=6.45,
    )
    add_image(
        doc,
        IMG_DIR_DESC / "category_pie_fusion_detail.png",
        "그림 4-2. 융복합 논문의 세부 법 분야 조합",
        width=6.45,
    )

    doc.add_heading("4.3 주요 키워드의 등장과 교체", level=2)
    add_body(
        doc,
        "연도별 상위 키워드의 변화는 법학계 AI 연구의 담론 중심이 이동했음을 보여준다. 2017년부터 2021년까지는 4차 산업혁명 키워드가 중심적 위치를 차지한 반면, 2023년 이후에는 생성형 인공지능이 가장 빠르게 부상하였다.",
    )
    add_image(
        doc,
        IMG_DIR_KEYWORD / "keyword_top10_by_year.png",
        "그림 4-3. 연도별 상위 10개 키워드 변화",
        width=6.7,
    )

    doc.add_heading("4.4 키워드 공출현 네트워크로 본 연구 주제군", level=2)
    add_body(
        doc,
        "이하에서는 저자 제공 키워드가 동일 논문에서 함께 등장한 관계를 바탕으로 AI 법학 연구의 주요 주제군을 살펴본다. 키워드 공출현 네트워크는 개별 키워드의 빈도만으로는 파악하기 어려운 연구 의제 간 결합 관계를 보여준다. 이 절의 관심은 네트워크 자체의 형식적 구조가 아니라, 어떤 법적 쟁점들이 하나의 연구 주제군으로 묶이는지에 있다.",
    )
    add_body(
        doc,
        "거대 컴포넌트에 Leiden 알고리즘을 적용한 결과, 키워드 공출현 네트워크는 14개 연구 주제군으로 분화되었다. 주요 주제군은 생성형 AI·저작권, 빅데이터·개인정보, 4차 산업혁명·거버넌스, 알고리즘·머신러닝, AI 기본법 등으로 요약된다.",
    )
    add_community_features_table(doc)

    add_body(
        doc,
        "또한 중심성 분석 결과에서도 4차 산업혁명, 빅데이터, 저작권, 생성형 인공지능 등이 높은 연결성을 보였는데, 이는 해당 키워드들이 분석 기간 동안 다양한 법적 쟁점과 반복적으로 결합되어 왔음을 시사한다.",
    )

    doc.add_heading("4.5 주요 연구 주제군의 시계열 변화", level=2)
    add_body(
        doc,
        "주요 연구 주제군의 누적 성장 양상은 연구 의제의 시기별 변화를 보여준다. 4차 산업혁명·거버넌스는 초기 포괄 담론을 대표하고, 빅데이터·개인정보와 알고리즘·머신러닝은 중기 이후 안정적으로 축적되며, 생성형 AI·저작권과 AI 기본법은 2023년 이후 빠르게 성장한다.",
    )
    add_image(
        doc,
        IMG_DIR_COMMUNITY / "kw_community_selected_growth.png",
        "그림 4-4. 주요 키워드 커뮤니티의 누적 논문 수 변화",
        width=6.75,
    )

    doc.add_heading("4.6 내부 인용의 증가와 연구 기반의 형성", level=2)
    add_body(
        doc,
        "인용 네트워크에 대해 커뮤니티 탐지를 시도하였으나, 스펙트럴 클러스터링 결과 전체 논문의 약 70%가 하나의 커뮤니티에 집중되는 편중이 확인되었다. 이는 인용 관계가 주제적 유사성뿐 아니라 기초문헌 및 일반 법리 문헌에 대한 공통 참조를 함께 반영하기 때문으로 볼 수 있다. 따라서 본 연구에서는 인용 네트워크의 커뮤니티 구조를 독립적인 주제 분류로 해석하기보다, 내부 인용 비율과 인용 시차 등 연구 축적성을 나타내는 지표를 중심으로 활용하였다.",
    )
    add_body(
        doc,
        "2016년 AI 법학 연구는 국내 선행연구의 축적이 아직 충분하지 않은 상태에서 출발하였다. 인공지능이 법학의 독립적 연구 주제로 본격적으로 다루어지기 시작한 시점이었기 때문에, 연구자들은 해외 문헌이나 인접 분야 논의에 상당 부분 의존할 수밖에 없었다. 실제로 2016년 AI 법학 논문의 내부 인용 비율은 8.3%에 그쳤고, 논문 한 편당 내부 선행연구 인용 건수도 평균 0.51건에 불과하였다.",
    )
    add_image(
        doc,
        IMG_DIR_CITATION / "07_internal_citation_flow.png",
        "그림 4-5. AI 법학 연구의 내부 인용 축적 추이",
        width=6.65,
    )
    add_body(
        doc,
        "그러나 2024년의 상황은 뚜렷하게 달라졌다. 내부 인용 비율은 40.5%로 높아졌고, 논문당 평균 내부 인용 건수는 4.17건으로 증가하였다. 내부 선행연구를 한 편 이상 인용한 논문의 비중도 2016년 31.9%에서 2024년 71.7%로 상승하였다. 이는 단순히 논문 수가 증가한 결과만은 아니다. 이 변화의 성격을 더 명확히 보여주는 지표는 인용 시차이다. 2016년 내부 인용의 평균 시차는 0.00년으로, 사실상 같은 해 발표된 논문 사이의 인용에 머물렀다. 반면 2024년에는 평균 시차가 3.11년으로 늘어났다. 이는 최근의 AI 법학 연구가 당해 연도의 동시적 논의만이 아니라, 수년에 걸쳐 축적된 국내 선행연구 위에서 전개되고 있음을 보여준다.",
    )
    add_body(
        doc,
        "이러한 수치들은 한국 AI 법학 연구가 하나의 담론장으로 성숙해가는 과정을 보여준다. 초기에는 각자의 문제의식에서 출발한 개별 연구들이 병렬적으로 등장했다면, 최근에는 선행연구를 의식하고 그 논의와 대화하면서 후속 논의를 쌓아가는 구조가 형성되고 있다. 키워드 공출현 분석이 연구 주제의 변화와 분화를 보여준다면, 인용 분석은 그 연구들이 시간의 흐름 속에서 얼마나 축적적 연구 기반을 형성해왔는지를 보조적으로 확인해준다.",
    )

    add_note(
        doc,
        "작성 메모: 본 문서는 4장 분석 결과 초안이며, 최종 논문 작성 시 표·그림 번호는 전체 논문 번호 체계에 맞추어 조정할 필요가 있다.",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
